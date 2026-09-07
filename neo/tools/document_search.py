from __future__ import annotations

import logging
import time
from pathlib import Path

from ..memory.document_index import DocumentIndex
from .base import RiskLevel, Tool, ToolResult

logger = logging.getLogger(__name__)

# Claude summarizes what comes back rather than this tool doing any NLP, so
# there is no point extracting more text than fits comfortably in one
# request; this is generous enough for most reports/spreadsheets while
# keeping token cost predictable.
MAX_CHARS = 20_000

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv"}

# Bounds for IndexFolderTool, same reasoning as find_file's search budget:
# a folder with thousands of documents shouldn't be able to make a single
# tool call run for minutes or index gigabytes of text into one row's
# worth of SQLite content per file.
MAX_FILES_PER_INDEX_RUN = 200
INDEX_TIME_BUDGET_SECONDS = 20.0


class UnsupportedDocumentError(Exception):
    """Raised when the file extension isn't one this tool knows how to read."""


class DocumentReadError(Exception):
    """Raised when the file exists and is a supported type, but reading it failed."""


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"--- {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _read_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


_READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
    ".txt": _read_plain_text,
    ".md": _read_plain_text,
    ".csv": _read_plain_text,
}


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    reader = _READERS.get(extension)
    if reader is None:
        raise UnsupportedDocumentError(
            f"'{extension}' uzantılı dosyaları okuyamıyorum. "
            f"Desteklenenler: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )
    try:
        text = reader(path)
    except Exception as exc:
        raise DocumentReadError(f"'{path.name}' okunurken bir hata oluştu.") from exc

    truncated = len(text) > MAX_CHARS
    return text[:MAX_CHARS] + ("\n\n[...metin kısaltıldı...]" if truncated else "")


class ReadDocumentTool(Tool):
    name = "read_document"
    description = (
        "Bir belgenin (PDF, Word .docx, Excel .xlsx, .txt/.md/.csv) metnini "
        "çıkarır ve sana gösterir, böylece içeriğini özetleyebilir ya da "
        "içinde arama yapabilirsin. Yol tam olarak bilinmiyorsa önce "
        "find_file ile dosyayı bul."
    )
    # Same privacy bar as capture_screen: reading a file the user points to
    # and sending its contents to the API is a one-time judgment call the
    # human should confirm, not something that should happen silently.
    risk = RiskLevel.MEDIUM
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Okunacak dosyanın tam yolu."},
        },
        "required": ["path"],
    }

    async def run(self, path: str, **kwargs: object) -> ToolResult:
        import asyncio

        file_path = Path(path).expanduser()
        if not file_path.is_file():
            return ToolResult(success=False, error=f"'{path}' diye bir dosya bulamadım.")

        try:
            text = await asyncio.to_thread(extract_text, file_path)
        except (UnsupportedDocumentError, DocumentReadError) as exc:
            return ToolResult(success=False, error=str(exc))

        return ToolResult(
            success=True,
            data={"path": str(file_path), "name": file_path.name, "text": text},
        )


def _index_folder_sync(index: DocumentIndex, folder: Path) -> tuple[int, int, bool]:
    """Runs off the asyncio loop -- walks the folder, extracts each
    supported file's text, and upserts it into the index. Returns
    (indexed_count, skipped_count, truncated)."""
    deadline = time.monotonic() + INDEX_TIME_BUDGET_SECONDS
    indexed = 0
    skipped = 0
    truncated = False
    for path in sorted(folder.rglob("*")):
        if indexed + skipped >= MAX_FILES_PER_INDEX_RUN or time.monotonic() > deadline:
            truncated = True
            break
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            text = extract_text(path)
        except (UnsupportedDocumentError, DocumentReadError):
            skipped += 1
            continue
        if not text.strip():
            skipped += 1
            continue
        index.upsert(str(path), path.name, text)
        indexed += 1
    return indexed, skipped, truncated


class IndexFolderTool(Tool):
    name = "index_folder"
    description = (
        "Bir klasördeki desteklenen belgeleri (PDF, Word, Excel, metin) "
        "aranabilir hale getirmek için içeriklerini yerel bir arama "
        "dizinine ekler. Bunu yaptıktan sonra search_documents ile bu "
        "klasördeki dosyalarda içerik araması yapılabilir. Büyük "
        "klasörlerde bir seferde en fazla "
        f"{MAX_FILES_PER_INDEX_RUN} dosya işlenir."
    )
    # Same bar as read_document: extracts and stores file contents locally
    # (nothing leaves the machine at index time -- only a later search's
    # matched snippet reaches the API), but still a one-time judgment call
    # about which of the user's folders NEO should be able to search.
    risk = RiskLevel.MEDIUM
    input_schema = {
        "type": "object",
        "properties": {
            "folder": {"type": "string", "description": "Dizine eklenecek klasörün tam yolu."},
        },
        "required": ["folder"],
    }

    def __init__(self, index: DocumentIndex) -> None:
        self._index = index

    async def run(self, folder: str, **kwargs: object) -> ToolResult:
        import asyncio

        folder_path = Path(folder).expanduser()
        if not folder_path.is_dir():
            return ToolResult(success=False, error=f"'{folder}' diye bir klasör bulamadım.")

        indexed, skipped, truncated = await asyncio.to_thread(
            _index_folder_sync, self._index, folder_path
        )
        return ToolResult(
            success=True,
            data={"indexed": indexed, "skipped": skipped, "truncated": truncated},
        )


class SearchDocumentsTool(Tool):
    name = "search_documents"
    description = (
        "Daha önce index_folder ile dizine eklenmiş belgeler arasında "
        "anahtar kelime/konu araması yapar (ör. 'yapay zekâ bağımsızlığı "
        "geçen ay hangi dosyada yazıyordu?'). En alakalı belgeleri, "
        "eşleşen kısmın kısa bir önizlemesiyle döner. Hiç belge dizine "
        "eklenmemişse boş sonuç döner -- önce index_folder çağır."
    )
    risk = RiskLevel.LOW
    input_schema = {
        "type": "object",
        "properties": {
            "query": {"type": "string", "description": "Aranacak konu ya da kelimeler."},
            "max_results": {"type": "integer", "description": "En fazla kaç sonuç (varsayılan 5)."},
        },
        "required": ["query"],
    }

    def __init__(self, index: DocumentIndex) -> None:
        self._index = index

    async def run(self, query: str, max_results: int = 5, **kwargs: object) -> ToolResult:
        import asyncio

        if not query.strip():
            return ToolResult(success=False, error="Ne arayacağımı belirtmen gerekiyor.")

        hits = await asyncio.to_thread(self._index.search, query, max_results)
        return ToolResult(
            success=True,
            data={
                "query": query,
                "count": len(hits),
                "results": [
                    {"path": h.path, "name": h.name, "score": round(h.score, 3), "snippet": h.snippet}
                    for h in hits
                ],
            },
        )
