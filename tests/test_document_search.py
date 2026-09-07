import asyncio

import pytest

from neo.tools.base import RiskLevel
from neo.tools.document_search import (
    DocumentReadError,
    MAX_CHARS,
    ReadDocumentTool,
    UnsupportedDocumentError,
    extract_text,
)


def test_extract_text_reads_plain_txt(tmp_path):
    path = tmp_path / "not.txt"
    path.write_text("merhaba dünya", encoding="utf-8")

    assert extract_text(path) == "merhaba dünya"


def test_extract_text_reads_markdown_and_csv(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("# başlık", encoding="utf-8")
    csv = tmp_path / "a.csv"
    csv.write_text("a,b\n1,2", encoding="utf-8")

    assert extract_text(md) == "# başlık"
    assert extract_text(csv) == "a,b\n1,2"


def test_extract_text_truncates_long_content(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("x" * (MAX_CHARS + 500), encoding="utf-8")

    result = extract_text(path)

    assert len(result) < MAX_CHARS + 500
    assert "kısaltıldı" in result


def test_extract_text_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "video.mp4"
    path.write_bytes(b"not really a video")

    with pytest.raises(UnsupportedDocumentError):
        extract_text(path)


def test_extract_text_wraps_reader_failures(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a real pdf")

    with pytest.raises(DocumentReadError):
        extract_text(path)


def test_extract_text_reads_xlsx(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "sheet.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sayfa1"
    sheet.append(["ad", "yaş"])
    sheet.append(["Ali", 30])
    workbook.save(path)

    result = extract_text(path)

    assert "Sayfa1" in result
    assert "Ali" in result


def test_extract_text_reads_docx(tmp_path):
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx/lxml not importable in this environment")
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("merhaba belge")
    document.save(path)

    result = extract_text(path)

    assert "merhaba belge" in result


def test_read_document_tool_reports_missing_file(tmp_path):
    result = asyncio.run(ReadDocumentTool().run(path=str(tmp_path / "yok.txt")))

    assert result.success is False


def test_read_document_tool_returns_text_for_existing_file(tmp_path):
    path = tmp_path / "not.txt"
    path.write_text("içerik burada", encoding="utf-8")

    result = asyncio.run(ReadDocumentTool().run(path=str(path)))

    assert result.success is True
    assert result.data["text"] == "içerik burada"
    assert result.data["name"] == "not.txt"


def test_read_document_tool_is_medium_risk():
    assert ReadDocumentTool.risk == RiskLevel.MEDIUM
